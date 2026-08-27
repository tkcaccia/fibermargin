#!/usr/bin/env python3
"""Apply a publication-oriented Word layout to Pandoc-generated FiberMargin files."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


INK = RGBColor(0x17, 0x22, 0x33)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x5F, 0x6B, 0x78)
LIGHT = "F4F6F9"

AUTHORS = (
    "Moussa Kassim; Martin Ocharo; Dalia Ahmed; Dupe Ojo; Alessia Vignoli; "
    "Leonardo Tenori; Dinesh Gupta; Silvano Piazza; Stefano Cacciatore"
)
AUTHOR_ROWS = (
    ("Moussa Kassim", "1,2", "Moussa.Kassim@icgeb.org"),
    ("Martin Ocharo", "1,2", "Martin.Ocharo@icgeb.org"),
    ("Dalia Ahmed", "1", "Dalia.Ahmed@icgeb.org"),
    ("Dupe Ojo", "1", "dupe.ojo@icgeb.org"),
    ("Alessia Vignoli", "3,4", "vignoli@cerm.unifi.it"),
    ("Leonardo Tenori", "3,4", "tenori@cerm.unifi.it"),
    ("Dinesh Gupta", "20", "dinesh.gupta@icgeb.org"),
    ("Silvano Piazza", "21,22", "piazza@icgeb.org"),
    ("Stefano Cacciatore", "1,2", "stefano.cacciatore@icgeb.org"),
)
AFFILIATIONS = (
    ("1", "Bioinformatics Unit, International Centre for Genetic Engineering and "
          "Biotechnology (ICGEB), Cape Town 7925, South Africa"),
    ("2", "Department of Integrative Biomedical Sciences, Institute of Infectious "
          "Disease & Molecular Medicine (IDM), University of Cape Town, Cape Town "
          "7925, South Africa"),
    ("3", "Department of Chemistry \"Ugo Schiff\", University of Florence, Sesto "
          "Fiorentino, Italy"),
    ("4", "Magnetic Resonance Center (CERM), University of Florence, Sesto "
          "Fiorentino, Italy"),
    ("20", "Translational Bioinformatics Group, International Centre for Genetic "
           "Engineering and Biotechnology (ICGEB), New Delhi, India"),
    ("21", "Computational Biology Group, International Centre for Genetic "
           "Engineering and Biotechnology (ICGEB), Trieste, Italy"),
    ("22", "Bioinformatics Facility, Department of Cellular, Computational and "
           "Integrative Biology - CIBIO, University of Trento, Trento, Italy"),
)
ACKNOWLEDGMENTS = (
    "The authors thank 10x Genomics and the maintainers of spatialLIBD, "
    "SpaGCN, and GraphST for making data or source code available for "
    "reproducible evaluation."
)
FUNDING = "The authors received no specific funding for this work."
COMPETING_INTERESTS = "The authors declare that they have no competing interests."
KEYWORDS = (
    "Keywords: categorical mask repair; spatial label refinement; coordinate-only "
    "learning; noisy labels; spatial transcriptomics"
)

def insert_before(reference: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addprevious(element)
    paragraph = Paragraph(element, reference._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_after(reference: Paragraph, text: str, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def set_cell_width(cell, width: float) -> None:
    width_twips = int(width * 1440)
    cell.width = Inches(width)
    properties = cell._tc.get_or_add_tcPr()
    width_element = properties.find(qn("w:tcW"))
    if width_element is None:
        width_element = OxmlElement("w:tcW")
        properties.append(width_element)
    width_element.set(qn("w:w"), str(width_twips))
    width_element.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table) -> None:
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_page_field(paragraph: Paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def set_font(style, name: str, size: float, color: RGBColor = INK, bold=None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_font(normal, "Calibri", 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name in ("Body Text", "First Paragraph"):
        if name in styles:
            set_font(styles[name], "Calibri", 11)
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            styles[name].paragraph_format.space_after = Pt(8)
            styles[name].paragraph_format.line_spacing = 1.25
            styles[name].paragraph_format.widow_control = True

    set_font(styles["Title"], "Calibri", 23, INK, True)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.space_after = Pt(10)
    styles["Title"].paragraph_format.keep_with_next = True

    if "Author" in styles:
        set_font(styles["Author"], "Calibri", 10.5, BLUE, True)
        styles["Author"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles["Author"].paragraph_format.space_after = Pt(4)

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, INK, 9, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        set_font(styles[name], "Calibri", size, color, True)
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.keep_together = True

    if "Compact" in styles:
        set_font(styles["Compact"], "Calibri", 10.5)
        styles["Compact"].paragraph_format.space_after = Pt(3)
        styles["Compact"].paragraph_format.line_spacing = 1.15

    if "Source Code" in styles:
        source_code = styles["Source Code"]
        set_font(source_code, "Consolas", 9)
        source_code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        source_code.paragraph_format.left_indent = Inches(0.18)
        source_code.paragraph_format.space_before = Pt(4)
        source_code.paragraph_format.space_after = Pt(0)
        source_code.paragraph_format.line_spacing = 1.0

    if "Verbatim Char" in styles:
        set_font(styles["Verbatim Char"], "Consolas", 9)

    if "Caption" in styles:
        set_font(styles["Caption"], "Calibri", 9.5, MUTED)
        styles["Caption"].font.italic = True
        styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles["Caption"].paragraph_format.space_before = Pt(4)
        styles["Caption"].paragraph_format.space_after = Pt(9)
        styles["Caption"].paragraph_format.keep_with_next = False

    if "Bibliography" in styles:
        set_font(styles["Bibliography"], "Calibri", 9.5)
        styles["Bibliography"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        styles["Bibliography"].paragraph_format.left_indent = Inches(0.25)
        styles["Bibliography"].paragraph_format.first_line_indent = Inches(-0.25)
        styles["Bibliography"].paragraph_format.space_after = Pt(5)
        styles["Bibliography"].paragraph_format.line_spacing = 1.05
        # References may break across pages; only the heading needs to stay with
        # the first entry. This avoids an unnecessary nearly blank final page.
        styles["Bibliography"].paragraph_format.keep_with_next = False
        styles["Bibliography"].paragraph_format.keep_together = False


def split_code_blocks(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        if paragraph.style.name != "Source Code":
            continue
        lines = paragraph.text.splitlines()
        if not lines:
            continue

        paragraph.text = lines[0]
        paragraph.style = "Source Code"
        reference = paragraph
        for line in lines[1:]:
            reference = insert_after(reference, line, "Source Code")
        reference.paragraph_format.space_after = Pt(8)


def replace_author_contact_block(author: Paragraph) -> None:
    """Format one author/email row per line, following the JMLR title-page style."""
    for child in list(author._p):
        if child.tag != qn("w:pPr"):
            author._p.remove(child)

    author.alignment = WD_ALIGN_PARAGRAPH.LEFT
    author.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT)
    author.paragraph_format.space_after = Pt(4)
    author.paragraph_format.line_spacing = 1.05
    author.paragraph_format.keep_with_next = True

    for row_index, (name, indices, email) in enumerate(AUTHOR_ROWS):
        name_run = author.add_run(name)
        name_run.bold = True
        name_run.font.name = "Calibri"
        name_run.font.size = Pt(10.5)
        name_run.font.color.rgb = INK

        index_run = author.add_run(indices)
        index_run.font.name = "Calibri"
        index_run.font.size = Pt(8)
        index_run.font.superscript = True
        index_run.font.color.rgb = INK

        author.add_run("\t")
        email_run = author.add_run(email)
        email_run.font.name = "Calibri"
        email_run.font.size = Pt(8.5)
        email_run.font.all_caps = True
        email_run.font.color.rgb = INK

        if row_index < len(AUTHOR_ROWS) - 1:
            email_run.add_break()


def populate_affiliations(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(14)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_with_next = True

    for row_index, (index, text) in enumerate(AFFILIATIONS):
        index_run = paragraph.add_run(index)
        index_run.font.name = "Calibri"
        index_run.font.size = Pt(8)
        index_run.font.superscript = True
        index_run.font.color.rgb = INK

        affiliation_run = paragraph.add_run(text)
        affiliation_run.font.name = "Calibri"
        affiliation_run.font.size = Pt(9)
        affiliation_run.font.italic = True
        affiliation_run.font.color.rgb = INK
        if row_index < len(AFFILIATIONS) - 1:
            affiliation_run.add_break()


def add_front_and_back_matter(document: Document, kind: str) -> None:
    author = next((p for p in document.paragraphs if p.style.name == "Author"), None)
    if author is not None:
        replace_author_contact_block(author)
        affiliations = insert_after(author, "", "Body Text")
        populate_affiliations(affiliations)

        if kind == "main" and not any(
            paragraph.text == "Abstract" for paragraph in document.paragraphs
        ):
            abstract = insert_after(affiliations, "Abstract", "Heading 1")
            abstract.paragraph_format.space_before = Pt(4)

    if kind == "main":
        introduction = next((p for p in document.paragraphs if p.text == "Introduction"), None)
        if introduction is not None:
            keywords = insert_before(introduction, KEYWORDS, "Body Text")
            keywords.paragraph_format.space_after = Pt(12)
            for run in keywords.runs:
                run.italic = True
                run.font.color.rgb = MUTED

    bibliography = next((p for p in document.paragraphs if p.style.name == "Bibliography"), None)
    if bibliography is not None:
        references = insert_before(bibliography, "References", "Heading 1")
        if kind == "main":
            existing_headings = {paragraph.text for paragraph in document.paragraphs}
            if "Acknowledgments" not in existing_headings:
                acknowledgments_text = insert_before(references, ACKNOWLEDGMENTS, "Body Text")
                acknowledgments_text.alignment = WD_ALIGN_PARAGRAPH.LEFT
                insert_before(acknowledgments_text, "Acknowledgments", "Heading 1")
            if "Funding" not in existing_headings:
                funding = insert_before(references, "", "Body Text")
                funding.alignment = WD_ALIGN_PARAGRAPH.LEFT
                funding.add_run("Funding. ").bold = True
                funding.add_run(FUNDING)
            if "Competing interests" not in existing_headings:
                competing = insert_before(references, "", "Body Text")
                competing.alignment = WD_ALIGN_PARAGRAPH.LEFT
                competing.add_run("Competing interests. ").bold = True
                competing.add_run(COMPETING_INTERESTS)


def number_headings(document: Document, kind: str) -> None:
    section = 0
    subsection = 0
    excluded = {"Abstract", "Acknowledgments", "References"}
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Heading 1" and paragraph.text not in excluded:
            section += 1
            subsection = 0
            prefix = f"S{section}" if kind == "supplement" else str(section)
            paragraph.text = f"{prefix}  {paragraph.text}"
        elif paragraph.style.name == "Heading 2":
            subsection += 1
            prefix = f"S{section}.{subsection}" if kind == "supplement" else f"{section}.{subsection}"
            paragraph.text = f"{prefix}  {paragraph.text}"


def configure_sections(document: Document, label: str) -> None:
    for index, section in enumerate(document.sections):
        section.start_type = WD_SECTION_START.NEW_PAGE if index else WD_SECTION_START.CONTINUOUS
        page_width = section.page_width
        page_height = section.page_height
        is_landscape = section.orientation == WD_ORIENT.LANDSCAPE or (
            page_width is not None and page_height is not None and page_width > page_height
        )
        if is_landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Inches(11)
            section.page_height = Inches(8.5)
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.55)
            section.left_margin = Inches(0.55)
            section.right_margin = Inches(0.55)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)
            section.top_margin = Inches(0.85)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        header = section.header
        header.is_linked_to_previous = False
        header_paragraph = header.paragraphs[0]
        header_paragraph.text = label
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header_paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.color.rgb = MUTED

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(footer_paragraph)
        for run in footer_paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.font.color.rgb = MUTED


def configure_tables(document: Document) -> None:
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for column_index, cell in enumerate(row.cells):
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0:
                    shade_cell(cell, LIGHT)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                    )
                    for run in paragraph.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(8)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = BLUE


def configure_figures(document: Document) -> None:
    max_width = Inches(6.6)
    for shape in document.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)
        paragraph = shape._inline
        while paragraph is not None and paragraph.tag != qn("w:p"):
            paragraph = paragraph.getparent()
        if paragraph is None:
            continue
        paragraph_properties = paragraph.get_or_add_pPr()
        justification = paragraph_properties.find(qn("w:jc"))
        if justification is None:
            justification = OxmlElement("w:jc")
            paragraph_properties.append(justification)
        justification.set(qn("w:val"), "center")


def number_figure_captions(document: Document) -> None:
    """Restore the Markdown figure numbers in the Word edition."""
    figure_number = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Image Caption":
            continue
        figure_number += 1
        prefix = f"Figure {figure_number}. "
        if not paragraph.text.startswith(prefix):
            # Change only the first run so embedded OMML remains intact.
            if paragraph.runs:
                paragraph.runs[0].text = prefix + paragraph.runs[0].text
            else:
                paragraph.add_run(prefix)
        paragraph.style = "Caption"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def insert_native_table_after(
    document: Document,
    anchor_text: str,
    caption_text: str,
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
) -> object:
    """Restore LaTeX-only manuscript tables in the Word edition."""
    anchor = next((p for p in document.paragraphs if anchor_text in p.text), None)
    if anchor is None:
        raise RuntimeError(f"Could not locate Word-table anchor: {anchor_text}")

    table = document.add_table(rows=1, cols=len(headers))
    for column, header in enumerate(headers):
        table.rows[0].cells[column].text = header
    for source_row in rows:
        row = table.add_row()
        for column, value in enumerate(source_row):
            row.cells[column].text = value

    caption = document.add_paragraph(caption_text, "Caption")
    caption.paragraph_format.keep_with_next = True
    anchor._p.addnext(caption._p)
    caption._p.addnext(table._tbl)
    return table


def configure_compact_result_table(table, widths: tuple[float, ...], font_size: float = 7.5) -> None:
    """Set fixed publication geometry for compact native Word result tables."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell, top=55, start=60, bottom=55, end=60)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.95
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(font_size if row_index else font_size - 0.4)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = BLUE
            if row_index == 0:
                shade_cell(cell, LIGHT)
    set_repeat_table_header(table.rows[0])


def insert_main_result_table(document: Document) -> object:
    return insert_native_table_after(
        document,
        "Table 1 places the primary real-coordinate matrices",
        "Table 1. Mean results across the predefined real-coordinate corruption "
        "units and frozen external prediction units. Best is the highest score "
        "among fixed controls applied in that panel. The Input column separates "
        "controlled corruptions from genuine saved upstream predictions. Sparse-region "
        "accuracy, reported in the supplement, uses the least prevalent reference "
        "class within an evaluation unit. Units derived from a common reference field "
        "are not treated as independent biological specimens.",
        ("Study", "Input", "Units", "Fiber\nAcc.", "Best\nAcc.", "Fiber\nARI", "Best\nARI"),
        (
            ("DLPFC", "Controlled", "45", "0.8942", "0.8846\n(GraphST)", "0.8220", "0.8112\n(Potts-like ICM)"),
            ("CRC", "Controlled", "60", "0.8589", "0.8604\n(Multiscale)", "0.7806", "0.7827\n(Multiscale)"),
            ("MERFISH", "Controlled", "45", "0.8905", "0.8843\n(alpha-Potts)", "0.8004", "0.7888\n(alpha-Potts)"),
            ("CamVid", "Saved\nprediction", "233", "0.6283", "0.6344\n(GraphST)", "0.5268", "0.5318\n(GraphST)"),
            ("S3DIS", "Saved\nprediction", "68", "0.6211", "0.6322\n(GraphST)", "0.5651", "0.5777\n(GraphST)"),
        ),
    )


def insert_figure_after(
    document: Document,
    anchor_text: str,
    image_path: Path,
    caption_text: str,
) -> None:
    anchor = next((p for p in document.paragraphs if anchor_text in p.text), None)
    if anchor is None:
        raise RuntimeError(f"Could not locate Word-figure anchor: {anchor_text}")
    if not image_path.exists():
        raise RuntimeError(f"The expected Word figure is missing: {image_path}")

    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(image_path), width=Inches(6.25))
    caption = document.add_paragraph(caption_text, "Caption")
    caption.paragraph_format.keep_with_next = True
    anchor._p.addnext(caption._p)
    anchor._p.addnext(figure._p)


def read_external_matrix(filename: str, columns: tuple[str, ...]) -> tuple[tuple[str, ...], ...]:
    root = Path(__file__).resolve().parents[1]
    path = root / "benchmarks" / "results" / "fibermargin_external_natural_validation_2026_07_20" / filename
    current_path = (
        root
        / "benchmarks"
        / "results"
        / "fibermargin_external_current_e373_v1"
        / "external_summary.csv"
    )
    if not path.exists():
        raise RuntimeError(f"The expected external result matrix is missing: {path}")
    if not current_path.exists():
        raise RuntimeError(f"The expected current FiberMargin summary is missing: {current_path}")

    dataset = "CamVid" if filename.startswith("camvid") else "S3DIS"
    with current_path.open(newline="", encoding="utf-8") as handle:
        current_rows = tuple(csv.DictReader(handle))
    current = next(
        (row for row in current_rows if row["dataset"] == dataset and row["rule"] == "FiberMargin"),
        None,
    )
    if current is None:
        raise RuntimeError(f"No current FiberMargin row found for {dataset}")

    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        source_rows = list(reader)
        initial = next(row for row in source_rows if row["method"] == "Initial")
        rows = []
        for row in source_rows:
            if row["method"] == "FiberMargin":
                for column in columns:
                    if column == "accuracy_gain":
                        row[column] = str(
                            float(current["accuracy"]) - float(initial["accuracy"])
                        )
                    elif column in current and current[column] not in ("", "NA"):
                        row[column] = current[column]
            values = [row["method"]]
            values.extend(f"{float(row[column]):.4f}" for column in columns)
            rows.append(tuple(values))
    return tuple(rows)


def insert_external_supplement_content(document: Document) -> tuple[object, object]:
    # Pandoc omits the raw-LaTeX external result tables from DOCX output, so
    # insert Word-native copies after the paragraph that introduces them.
    anchor = next(
        (
            p
            for p in document.paragraphs
            if "Table S9 reports all CamVid outcomes" in p.text
        ),
        None,
    )
    if anchor is None:
        raise RuntimeError("Could not locate external-supplement result-table anchor")

    camvid_table = document.add_table(rows=1, cols=6)
    for column, header in enumerate(("Method", "Accuracy", "Gain", "Mean IoU", "ARI", "Damage")):
        camvid_table.rows[0].cells[column].text = header
    for source_row in read_external_matrix(
        "camvid_complete_matrix.csv",
        ("accuracy", "accuracy_gain", "mean_iou", "ari", "damage_rate"),
    ):
        row = camvid_table.add_row()
        for column, value in enumerate(source_row):
            row.cells[column].text = value
    camvid_caption = document.add_paragraph(
        "Table S9. CamVid external upstream-prediction matrix across 233 official test images. Gain is accuracy gain over the saved upstream categorical prediction. Higher is better except for damage.",
        "Caption",
    )

    s3dis_table = document.add_table(rows=1, cols=7)
    for column, header in enumerate(("Method", "Accuracy", "Gain", "ARI", "Boundary", "Sparse", "Damage")):
        s3dis_table.rows[0].cells[column].text = header
    for source_row in read_external_matrix(
        "s3dis_complete_matrix.csv",
        ("accuracy", "accuracy_gain", "ari", "boundary_accuracy", "sparse_region_accuracy", "damage_rate"),
    ):
        row = s3dis_table.add_row()
        for column, value in enumerate(source_row):
            row.cells[column].text = value
    s3dis_caption = document.add_paragraph(
        "Table S10. S3DIS external upstream-prediction matrix across 68 held-out Area 5 room blocks. Boundary accuracy uses eight Euclidean nearest neighbours; sparse accuracy is the least frequent reference class within a block. Higher is better except for damage.",
        "Caption",
    )

    for block in reversed((
        # Keep the second caption with its table in Word.  The PDF preserves the
        # manuscript's bottom-caption convention through its native LaTex table.
        camvid_table._tbl, camvid_caption._p, s3dis_caption._p, s3dis_table._tbl,
    )):
        anchor._p.addnext(block)
    return camvid_table, s3dis_table



def build(input_path: Path, output_path: Path, kind: str) -> None:
    document = Document(input_path)
    document.core_properties.author = AUTHORS
    document.core_properties.subject = "Multiclass spatial label repair"
    document.core_properties.keywords = (
        "categorical masks; spatial labels; FiberMargin; selective prediction"
    )

    configure_styles(document)
    split_code_blocks(document)
    add_front_and_back_matter(document, kind)
    number_headings(document, kind)
    main_results = None
    supplement_results = None
    if kind == "main":
        # The current manuscript carries its complete native tables and figures.
        # Earlier revisions injected an extra result table and a hard-coded
        # Figure 3 here, which duplicated and misnumbered the present content.
        pass
    elif kind == "supplement":
        supplement_results = insert_external_supplement_content(document)
    if kind == "main":
        label = "FiberMargin | JMLR manuscript"
    elif kind == "methods":
        label = "FiberMargin | Materials and Methods"
    else:
        label = "FiberMargin | Supplement"
    configure_sections(document, label)
    configure_tables(document)
    if main_results is not None:
        configure_compact_result_table(
            main_results, (0.75, 0.98, 0.40, 0.65, 1.18, 0.65, 1.25), font_size=6.8
        )
    if supplement_results is not None:
        configure_compact_result_table(
            supplement_results[0], (1.70, 0.68, 0.52, 0.78, 0.52, 0.58), font_size=7.2
        )
        configure_compact_result_table(
            supplement_results[1], (1.70, 0.63, 0.50, 0.52, 0.68, 0.60, 0.55), font_size=7.0
        )
    configure_figures(document)
    if kind == "main":
        number_figure_captions(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--kind", choices=("main", "methods", "supplement"), required=True
    )
    args = parser.parse_args()
    build(args.input, args.output, args.kind)


if __name__ == "__main__":
    main()
